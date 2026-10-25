import io
import pytest
from datetime import date, timedelta
from fastapi.testclient import TestClient
import docx
import pypdf

from app.main import app
from app.db import Base, engine, get_db
from app.seed import seed_database
from app.models import User, Application, Document, ScrutinyFinding, CorrectionRequest


@pytest.fixture(scope="module", autouse=True)
def setup_test_db():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with next(get_db()) as db:
        seed_database(db)


@pytest.fixture
def client():
    return TestClient(app)


def get_token(client, email, password="DemoPass123!"):
    res = client.post("/api/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, res.text
    return res.json()["access_token"]


def create_sample_pdf(text: str) -> bytes:
    """Helper to generate in-memory PDF with pypdf."""
    from pypdf import PdfWriter
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    # Using simple BytesIO
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_sample_docx(text: str) -> bytes:
    """Helper to generate in-memory DOCX with python-docx."""
    doc = docx.Document()
    doc.add_heading("AICTE Institutional Regulatory Document", 0)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- 1. Institution Application Lifecycle ---

def test_institution_create_and_patch_application(client):
    inst_token = get_token(client, "college@demo.local")
    
    # 1. Create Application
    res = client.post(
        "/api/v1/applications",
        json={
            "application_type": "Extension of Approval",
            "notes": "Draft EoA application 2026-2027",
            "academic_year": "2026-2027",
            "form_data": {
                "institution_code": "ABC-ENG-001",
                "total_intake": 480,
                "programmes": ["Computer Science", "Electronics", "Mechanical"]
            }
        },
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert res.status_code == 201
    data = res.json()
    app_id = data["id"]
    assert data["institution_name"] == "ABC Engineering College"
    assert "AICTE-" in data["reference"]

    # 2. Patch Application Form Data
    patch_res = client.patch(
        f"/api/v1/applications/{app_id}",
        json={
            "notes": "Updated EoA notes after committee review",
            "form_data": {
                "institution_code": "ABC-ENG-001",
                "total_intake": 540,
                "programmes": ["Computer Science", "Electronics", "Mechanical", "AI & Data Science"]
            }
        },
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert patch_res.status_code == 200
    patched_data = patch_res.json()
    assert "Updated EoA notes" in patched_data["notes"]
    assert "540" in patched_data["form_data"]


def test_institution_document_upload_and_extraction(client):
    inst_token = get_token(client, "college@demo.local")
    
    # Create sample docx document
    docx_bytes = create_sample_docx(
        "Registered Land Lease Deed between Lessor Alpha Trust and Lessee ABC Engineering College. "
        "Lease tenure: 35 years. Khasra No. 504, measuring 5.2 Acres. Registered with Sub-Registrar."
    )
    
    # Upload to application 1
    res = client.post(
        "/api/v1/applications/1/documents",
        files={"file": ("lease_deed_35yr.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"document_type": "LAND_OWNERSHIP_DEED"},
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert res.status_code == 200
    data = res.json()
    assert "document" in data
    assert data["document"]["document_type"] == "LAND_OWNERSHIP_DEED"
    assert "ai_analysis" in data
    assert data["ai_analysis"]["extracted_fields"]["Lease Tenure"] == "35 Years"


def test_application_timeline_and_notifications(client):
    inst_token = get_token(client, "college@demo.local")
    
    # 1. Fetch Timeline
    res = client.get("/api/v1/applications/1/timeline", headers={"Authorization": f"Bearer {inst_token}"})
    assert res.status_code == 200
    timeline = res.json()
    assert len(timeline) >= 1
    actions = [t["action"] for t in timeline]
    assert "DOC_UPLOADED" in actions or "SUBMITTED" in actions

    # 2. Fetch Notifications
    notif_res = client.get("/api/v1/notifications", headers={"Authorization": f"Bearer {inst_token}"})
    assert notif_res.status_code == 200
    notifs = notif_res.json()
    assert len(notifs) >= 1

    # 3. Mark Notification Read
    notif_id = notifs[0]["id"]
    read_res = client.patch(f"/api/v1/notifications/{notif_id}/read", headers={"Authorization": f"Bearer {inst_token}"})
    assert read_res.status_code == 200
    assert read_res.json()["is_read"] is True


# --- 2. Officer-Facing Scrutiny & Workflow Controls ---

def test_officer_rbac_protection(client):
    inst_token = get_token(client, "college@demo.local")
    # Institution user must be denied access to officer endpoints
    res = client.get("/api/v1/officer/applications", headers={"Authorization": f"Bearer {inst_token}"})
    assert res.status_code == 403


def test_officer_list_and_inspect_applications(client):
    officer_token = get_token(client, "officer@demo.local")
    res = client.get("/api/v1/officer/applications", headers={"Authorization": f"Bearer {officer_token}"})
    assert res.status_code == 200
    apps = res.json()
    assert len(apps) >= 1

    # Detailed view
    app_id = apps[0]["id"]
    det_res = client.get(f"/api/v1/officer/applications/{app_id}", headers={"Authorization": f"Bearer {officer_token}"})
    assert det_res.status_code == 200
    details = det_res.json()
    assert "application" in details
    assert "documents" in details
    assert "audit_history" in details


def test_officer_run_scrutiny_pipeline(client):
    officer_token = get_token(client, "officer@demo.local")
    res = client.post(
        "/api/v1/officer/applications/1/scrutiny",
        json={"run_type": "INITIAL"},
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert res.status_code == 200
    data = res.json()
    assert "scrutiny_run" in data
    assert "findings" in data
    assert "deficiency_count" in data
    assert len(data["findings"]) >= 1


def test_end_to_end_correction_request_and_resubmission(client):
    officer_token = get_token(client, "officer@demo.local")
    inst_token = get_token(client, "college@demo.local")

    # 1. Officer requests corrections
    req_res = client.post(
        "/api/v1/officer/applications/1/request-corrections",
        json={
            "reason": "Fire Safety NOC has expired and Faculty cadre ratio is incomplete.",
            "deficiencies": [
                {"rule": "APH-2024-CH2-FIRE", "description": "Fire NOC validity expired."},
                {"rule": "APH-2024-CH2-FACULTY", "description": "Cadre ratio short by 2 Associate Professors."}
            ]
        },
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert req_res.status_code == 200
    app_data = req_res.json()["application"]
    assert app_data["current_stage"] == "RETURNED_FOR_CLARIFICATION"
    assert app_data["status"] == "RETURNED"

    # 2. Institution resubmits with clarification
    resubmit_res = client.post(
        "/api/v1/applications/1/resubmit",
        json={
            "response_text": "Updated Fire NOC renewed by District Authority attached. Appointed 2 new Associate Professors."
        },
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert resubmit_res.status_code == 200
    resubmitted_app = resubmit_res.json()
    assert resubmitted_app["current_stage"] == "RE_SCRUTINY"
    assert resubmitted_app["status"] == "IN_PROGRESS"

    # 3. Officer completes re-scrutiny
    complete_res = client.post(
        "/api/v1/officer/applications/1/complete-rescrutiny",
        json={
            "decision": "CLEARED",
            "remarks": "Verified renewed Fire NOC and appointment orders. All deficiencies cleared."
        },
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert complete_res.status_code == 200
    cleared_app = complete_res.json()
    # Route for application 1 is Extension of Approval -> goes to COMPLIANCE_PENDING
    assert cleared_app["current_stage"] in ["COMPLIANCE_PENDING", "EVC_PENDING"]
    assert cleared_app["deficiency_count"] == 0


def test_officer_statutory_decision(client):
    officer_token = get_token(client, "officer@demo.local")

    # Issue statutory decision (APPROVED)
    res = client.post(
        "/api/v1/officer/applications/1/decision",
        json={
            "decision": "APPROVED",
            "remarks": "Extension of Approval granted for academic year 2026-2027.",
            "approval_order_reference": "AICTE/EOA/2026/NORTH-001"
        },
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "COMPLETED"
    assert data["decision_outcome"] == "APPROVED"
    assert data["decision_by"] == "officer@demo.local"


# --- 3. Route 2: New Programme / Level Workflow ---

def test_new_programme_mandatory_evc_route(client):
    inst_token = get_token(client, "college@demo.local")
    officer_token = get_token(client, "officer@demo.local")

    # 1. Create New Programme Application
    create_res = client.post(
        "/api/v1/applications",
        json={
            "application_type": "New Programme / Level",
            "notes": "Starting new Master in Computer Applications (MCA) programme.",
            "academic_year": "2026-2027",
            "form_data": {
                "programme": "MCA",
                "proposed_intake": 60,
                "building_block": "Block D - Computing Science"
            }
        },
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert create_res.status_code == 201
    new_app_id = create_res.json()["id"]

    # 2. Run Scrutiny on New Programme
    scrutiny_res = client.post(
        f"/api/v1/officer/applications/{new_app_id}/scrutiny",
        json={"run_type": "INITIAL"},
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert scrutiny_res.status_code == 200
    scrutiny_data = scrutiny_res.json()
    # EVC procedural requirement must be included for New Programme
    rules = [f["rule_id"] for f in scrutiny_data["findings"]]
    assert any("EVC" in r for r in rules)

    # 3. Advance to EVC_PENDING
    adv_res = client.post(
        f"/api/v1/officer/applications/{new_app_id}/advance",
        json={
            "target_stage": "EVC_PENDING",
            "reason": "Primary scrutiny cleared; scheduling mandatory EVC physical inspection.",
            "assigned_to": "EVC Coordinator"
        },
        headers={"Authorization": f"Bearer {officer_token}"}
    )
    assert adv_res.status_code == 200
    assert adv_res.json()["current_stage"] == "EVC_PENDING"


# --- 4. Insecure Direct Object Reference (IDOR) & Security Safeguards ---

def test_idor_protection_between_institutions(client):
    inst_token = get_token(client, "college@demo.local")
    
    # Application 2 belongs to Sunrise Institute of Technology
    # GET detail should return 403 Forbidden
    res = client.get("/api/v1/applications/2", headers={"Authorization": f"Bearer {inst_token}"})
    assert res.status_code == 403

    # PATCH should return 403 Forbidden
    patch_res = client.patch(
        "/api/v1/applications/2",
        json={"notes": "Hacked cross-tenant modification"},
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert patch_res.status_code == 403

    # Resubmit should return 403 Forbidden
    resub_res = client.post(
        "/api/v1/applications/2/resubmit",
        json={"response_text": "Unauthorized resubmit"},
        headers={"Authorization": f"Bearer {inst_token}"}
    )
    assert resub_res.status_code == 403
